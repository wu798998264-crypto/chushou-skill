#!/usr/bin/env python3
"""Offline, auditable Chinese drama transcription document builder.

Usage: python build_drama_documents.py --manifest manifest.json --out-dir output
       [--docx] [--require-complete]

No downloads, recognition, model calls, or network access are performed. Paths
inside the manifest resolve against its directory, never the working directory.
Exit codes: 0 fully reviewed; 2 missing/unreviewed episodes; 1 invalid evidence
or an output failure. Existing differing outputs are saved as numbered versions.
"""
import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import unicodedata

from compile_drama_transcripts import write_versioned


def read_json(path):
    return json.loads(Path(path).read_bytes().decode("utf-8-sig"))


def object_value(value, label):
    if not isinstance(value, dict):
        raise ValueError(label + " 必须是 JSON 对象")
    return value


def text_value(value, label, *, empty=False):
    if not isinstance(value, str) or (not empty and not value.strip()):
        raise ValueError(label + " 必须是" + ("字符串" if empty else "非空字符串"))
    value.encode("utf-8", errors="strict")
    return value


def positive_integer(value, label):
    if type(value) is not int or value < 1:
        raise ValueError(label + " 必须是正整数")
    return value


def sha_value(value, label):
    if not isinstance(value, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", value):
        raise ValueError(label + " 必须是 64 位 SHA256")
    return value.lower()


def resolved(value, base, label):
    path = Path(text_value(value, label)).expanduser()
    return (path if path.is_absolute() else base / path).resolve()


def file_sha(path):
    result = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def verify_file(path, expected, label):
    expected = sha_value(expected, label + " SHA256")
    if not path.is_file():
        raise ValueError(label + " 文件不存在：" + str(path))
    actual = file_sha(path)
    if actual != expected:
        raise ValueError(label + " SHA256 不匹配：" + str(path))
    return actual


def episode_index(values, total, label, series_id=None):
    if not isinstance(values, list):
        raise ValueError(label + " episodes 必须是数组")
    result = {}
    for item in values:
        object_value(item, label + " episode")
        number = positive_integer(item.get("episode"), label + " episode")
        if number > total or number in result:
            raise ValueError(label + " 存在重复或越界集次：" + str(number))
        if "series_id" in item and item["series_id"] != series_id:
            raise ValueError(label + " 包含其他剧集：第" + str(number) + "集")
        result[number] = item
    return result


def selected_numbers(value, total, label):
    if not isinstance(value, list) or not value:
        raise ValueError(label + " 必须是非空集次数组")
    result = []
    for number in value:
        positive_integer(number, label + " episode")
        if number > total or number in result:
            raise ValueError(label + " 存在重复或越界集次")
        result.append(number)
    return sorted(result)


def load_download_binding(path, manifest):
    """Validate the entire source catalog, not only the selected episodes."""
    batch = object_value(read_json(path), "下载清单")
    series_id = text_value(manifest.get("series_id"), "manifest series_id")
    if batch.get("series_id") != series_id:
        raise ValueError("下载清单与 manifest 的 series_id 不匹配")
    source_total = batch.get("source_total_episodes", batch.get("source_total"))
    if positive_integer(source_total, "下载清单 source_total_episodes") != manifest["total_episodes"]:
        raise ValueError("下载清单与 manifest 的总集数不匹配")
    if "source_total" in batch and batch["source_total"] != source_total:
        raise ValueError("下载清单两个 source_total 字段冲突")
    selected = manifest.get("selected_episodes", list(range(1, manifest["total_episodes"] + 1)))
    batch_selection = selected_numbers(batch.get("selected_episodes", list(range(1, source_total + 1))), source_total, "下载清单 selected_episodes")
    if not set(selected).issubset(batch_selection):
        raise ValueError("manifest 的 selected_episodes 选集范围超出下载批次范围")
    catalog_path = resolved(batch.get("catalog"), path.parent, "下载清单 catalog")
    catalog = object_value(read_json(catalog_path), "源目录")
    if catalog.get("series_id") != series_id:
        raise ValueError("源目录与 manifest 的 series_id 不匹配")
    total = positive_integer(catalog.get("total_episodes"), "源目录 total_episodes")
    if total != manifest["total_episodes"]:
        raise ValueError("源目录与 manifest 的总集数不匹配")
    items = episode_index(catalog.get("episodes"), total, "源目录", series_id)
    if set(items) != set(range(1, total + 1)):
        raise ValueError("源目录不是 1 至 total_episodes 的真实全集目录")
    seen_ids = set()
    for item in items.values():
        video_id = text_value(item.get("video_id"), "源目录 video_id")
        if video_id in seen_ids:
            raise ValueError("源目录 video_id 重复，无法确定集次")
        seen_ids.add(video_id)
    downloads = episode_index(batch.get("episodes"), total, "下载清单", series_id)
    if set(downloads) - set(batch_selection):
        raise ValueError("下载清单 episodes 包含其 selected_episodes 范围外的集次")
    for number, item in downloads.items():
        if item.get("video_id") != items[number]["video_id"]:
            raise ValueError("下载清单 video_id 与源目录不匹配：第" + str(number) + "集")
        if item.get("status") in {"downloaded", "reused"}:
            sha_value(item.get("sha256"), "下载清单 sha256")
        else:
            text_value(item.get("status"), "下载清单 status")
    return items, downloads, {
        "path": str(path), "sha256": file_sha(path), "catalog": str(catalog_path),
        "catalog_sha256": file_sha(catalog_path), "series_id": series_id,
        "source_total_episodes": total, "selected_episodes": batch_selection,
    }


def load_manifest(path):
    path = Path(path).resolve()
    data = object_value(read_json(path), "manifest")
    if type(data.get("schema_version")) is not int or data["schema_version"] != 1:
        raise ValueError("manifest schema_version 必须为整数 1")
    text_value(data.get("title"), "title")
    total = positive_integer(data.get("total_episodes"), "total_episodes")
    if "series_id" in data:
        text_value(data["series_id"], "series_id")
    entries = episode_index(data.get("episodes"), total, "manifest", data.get("series_id"))
    data["selected_episodes"] = selected_numbers(data.get("selected_episodes", list(range(1, total + 1))), total, "selected_episodes")
    if set(entries) - set(data["selected_episodes"]):
        raise ValueError("manifest episodes 包含所选范围外的集次")
    catalogs, downloads, binding = None, None, None
    if data.get("download_manifest") is not None:
        batch_path = resolved(data["download_manifest"], path.parent, "download_manifest")
        catalogs, downloads, binding = load_download_binding(batch_path, data)
    return data, entries, catalogs, downloads, binding


def semantic_characters(text):
    """Ignore common sentence punctuation, protecting semantic punctuation.

    Unicode P* alone is too broad: %, /, - and decimal points are punctuation
    yet can change the meaning of numbers, units, words and URLs. Unknown marks
    are retained conservatively. Whitespace is ignored before context checks.
    This is a loss-detection guard, not a semantic-equivalence proofreader.
    """
    compact = "".join(char for char in text if not char.isspace())
    # Protect URL punctuation including query separators and fragment markers.
    # Sentence marks in URLs are deliberately treated conservatively; a caller
    # can make a documented reviewed edit when the boundary is ambiguous.
    protected = set()
    for match in re.finditer(r"(?:https?://|www\.)[^\s<>「」『』“”‘’，。！？；、]+", text, re.I):
        start = len("".join(char for char in text[:match.start()] if not char.isspace()))
        length = len("".join(char for char in match.group() if not char.isspace()))
        protected.update(range(start, start + length))
    sentence_marks = set('。，、；：！？…“”‘’「」『』《》〈〉,;:.!?"\'')
    result = []
    for index, char in enumerate(compact):
        previous = compact[index - 1] if index else ""
        following = compact[index + 1] if index + 1 < len(compact) else ""
        # Decimal/grouping marks and time/ratio separators inside numbers.
        numeric_mark = char in ".．,，:：" and previous.isdigit() and following.isdigit()
        # Dot/colon in Latin names/units and apostrophes inside words.
        word_mark = char in ".:'’" and previous.isascii() and previous.isalnum() and following.isascii() and following.isalnum()
        if (char in sentence_marks and unicodedata.category(char).startswith("P")
                and index not in protected and not numeric_mark and not word_mark):
            continue
        result.append(char)
    return "".join(result)


def require_punctuation_only(before, after, label):
    if semantic_characters(before) != semantic_characters(after):
        raise ValueError(label + " 改变了非标点正文或语义标记（含汉字、字母、数字、小数/正负号/百分号/单位/URL）")


def apply_review(raw, raw_sha, path):
    """Apply non-overlapping original Unicode code-point offsets exactly once."""
    if path is None:
        return raw, {"reviewed": False, "edits": [], "uncertain": [], "reason": "未提供校对文件"}
    review = object_value(read_json(path), "review_file")
    if sha_value(review.get("raw_sha256"), "review raw_sha256") != raw_sha:
        raise ValueError("review raw_sha256 与原始识别文件不匹配")
    if "reviewed" in review and type(review["reviewed"]) is not bool:
        raise ValueError("reviewed 必须是布尔值")
    edits = review.get("edits", [])
    uncertain = review.get("uncertain", [])
    if not isinstance(edits, list) or not isinstance(uncertain, list):
        raise ValueError("review edits 和 uncertain 必须是数组")
    audited = []
    for edit in edits:
        object_value(edit, "edit")
        start = edit.get("start")
        if type(start) is not int or start < 0 or start > len(raw):
            raise ValueError("edit start 必须是原文 Unicode 码点范围内的整数")
        before = text_value(edit.get("before"), "edit before", empty=True)
        after = text_value(edit.get("after"), "edit after", empty=True)
        text_value(edit.get("reason"), "edit reason")
        evidence = edit.get("evidence")
        if not ((isinstance(evidence, str) and evidence.strip()) or (isinstance(evidence, dict) and evidence)):
            raise ValueError("edit evidence 必须是非空字符串或证据对象")
        if before == after:
            raise ValueError("edit before 和 after 相同，不是有效校订")
        if raw[start:start + len(before)] != before:
            raise ValueError("edit before 与原始偏移不精确匹配：" + str(start))
        audited.append({**edit, "end": start + len(before)})
    audited.sort(key=lambda item: item["start"])
    previous_end, previous_start = -1, -1
    for edit in audited:
        if edit["start"] < previous_end or edit["start"] == previous_start:
            raise ValueError("edits 在原文中的范围重叠或起点重复")
        previous_start, previous_end = edit["start"], edit["end"]
    for item in uncertain:
        object_value(item, "uncertain item")
        text_value(item.get("text"), "uncertain text")
        text_value(item.get("reason"), "uncertain reason")
    cursor, chunks = 0, []
    for edit in audited:
        chunks.extend((raw[cursor:edit["start"]], edit["after"]))
        cursor = edit["end"]
    chunks.append(raw[cursor:])
    return "".join(chunks), {
        "reviewed": review.get("reviewed") is True, "edits": audited, "uncertain": uncertain,
        "review_file": str(path), "review_sha256": file_sha(path),
    }


def sentence_lines(text):
    """Put sentence endings on separate lines without deleting semantic text."""
    normalized = re.sub(r"\s+", " ", text).strip()
    output, pending = [], False
    closers = set('”’」』）】》〉"\'')
    for index, char in enumerate(normalized):
        terminal = char in "。！？!?…" or (char == "." and not (
            index > 0 and index + 1 < len(normalized) and normalized[index - 1].isdigit() and normalized[index + 1].isdigit()))
        if pending and char not in closers and not terminal:
            if output and output[-1] != "\n":
                output.append("\n")
            pending = False
        if char == " " and output and output[-1] == "\n":
            continue
        output.append(char)
        pending = pending or terminal
    formatted = "".join(output).strip()
    require_punctuation_only(text, formatted, "一句一行排版")
    return formatted


def process_episode(number, item, base, catalog=None, downloads=None):
    report = {"episode": number, "status": "missing", "reviewed": False, "edits": [], "uncertain": []}
    if item is None:
        report["reason"] = "未列入处理清单，缺少本集识别结果"
        return report, None, None, None
    try:
        if catalog is not None:
            if item.get("video_id") != catalog[number]["video_id"]:
                raise ValueError("manifest video_id 与源目录不匹配")
            download = downloads.get(number)
            if not download or download.get("status") not in {"downloaded", "reused"}:
                raise ValueError("本集下载未完成或缺少下载记录")
            if object_value(download.get("metadata", {}), "下载 metadata").get("has_audio") is False:
                raise ValueError("本集已确认没有音轨")
            if sha_value(item.get("source_sha256"), "source_sha256") != sha_value(download.get("sha256"), "下载 sha256"):
                raise ValueError("source_sha256 与已下载结果不匹配")
            report["download_status"] = download["status"]
        if item.get("video_id") is not None:
            report["video_id"] = text_value(item["video_id"], "video_id")
        source = resolved(item.get("source_file"), base, "source_file")
        raw_path = resolved(item.get("raw_file"), base, "raw_file")
        report.update({"source_file": str(source), "raw_file": str(raw_path)})
        backend = object_value(item.get("backend"), "backend")
        report["backend"] = {key: text_value(backend.get(key), "backend " + key) for key in ("name", "url", "model")}
        for key, path, label in (("source", source, "源视频"), ("raw", raw_path, "原始识别")):
            expected = sha_value(item.get(key + "_sha256"), label + " SHA256")
            report[key + "_sha256"] = expected
            report[key + "_sha256_verified"] = False
            if not path.is_file():
                raise ValueError(label + " 文件不存在：" + str(path))
            actual = file_sha(path)
            report[key + "_sha256_actual"] = actual
            if actual != expected:
                raise ValueError(label + " SHA256 不匹配：" + str(path))
            report[key + "_sha256_verified"] = True
        # Decode bytes directly: preserve original line endings, BOM and all
        # content. There is no universal-newline or punctuation normalization.
        raw = raw_path.read_bytes().decode("utf-8", errors="strict")
        if not raw.lstrip("\ufeff").strip():
            raise ValueError("原始识别正文为空")
        review_path = resolved(item["review_file"], base, "review_file") if item.get("review_file") is not None else None
        corrected, audit = apply_review(raw, report["raw_sha256"], review_path)
        report.update(audit)
        if item.get("punct_file") is not None:
            punct_path = resolved(item["punct_file"], base, "punct_file")
            punct = punct_path.read_bytes().decode("utf-8", errors="strict")
            require_punctuation_only(corrected, punct, "punct_file")
            report["punctuation"] = {"file": str(punct_path), "sha256": file_sha(punct_path)}
        else:
            punct = corrected
        formatted = sentence_lines(punct)
        require_punctuation_only(corrected, formatted, "最终排版")
        report.update({"status": "reviewed" if report["reviewed"] else "pending_review",
                       "raw_characters": len(raw), "corrected_characters": len(corrected),
                       "formatted_characters": len(formatted),
                       "corrected_sha256": hashlib.sha256(corrected.encode("utf-8")).hexdigest(),
                       "formatted_sha256": hashlib.sha256(formatted.encode("utf-8")).hexdigest()})
        return report, raw, corrected, formatted
    except (OSError, ValueError, TypeError, KeyError) as exc:
        report.update({"status": "failed", "reviewed": False, "reason": str(exc)})
        return report, None, None, None


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def safe_name(title):
    result = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).rstrip(" .")[:100].rstrip(" .") or "短剧"
    if result.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *["COM" + str(i) for i in range(1, 10)], *["LPT" + str(i) for i in range(1, 10)]}:
        result = "_" + result
    return result


def docx_bytes(text, title):
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("--docx 需要可选依赖 python-docx；TXT 和审计不依赖它") from exc
    document = Document()
    document.core_properties.title = title
    document.core_properties.subject = "按集整理的音频识别及校对文稿"
    for line in text.splitlines():
        document.add_paragraph(line)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_document(title, series_id, total, selected, episodes, sections, raw=False, complete=False):
    full_series = selected == list(range(1, total + 1))
    label = ("全集原始识别全文" if full_series else "选集原始识别全文") if raw else ("全集" if full_series else "选集") + ("校对版" if complete else "校对草稿")
    header = [title + " · " + label]
    if series_id:
        header.append("剧集编号：" + series_id)
    header.append("总集数：" + str(total))
    header.append("处理范围：" + ("全集" if full_series else "选集") + "（第" + "、".join(str(number) for number in selected) + "集）")
    header.append("原始机器识别，未作内容改写。" if raw else "校订有逐项证据；未确认内容保留待核标记，详见审计报告。")
    result = ["\n".join(header)]
    for entry, body in zip(episodes, sections):
        note = ""
        if body is None:
            body = "【未完成：" + entry.get("reason", "缺少结果") + "】"
        elif not raw:
            if not entry["reviewed"]:
                note = "【待校对】\n"
            if entry["uncertain"]:
                note += "【待核事项：" + str(len(entry["uncertain"])) + "项，详见审计报告】\n"
        result.append("第" + str(entry["episode"]) + "集\n\n" + note + body)
    return "\n\n".join(result) + "\n"


def build(manifest_path, out_dir, *, docx=False, require_complete=False):
    """Build outputs and always retain a machine report on validation failure."""
    manifest_path, out_dir = Path(manifest_path).resolve(), Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    report = {"schema_version": 1, "manifest": str(manifest_path), "complete": False,
              "recognition_complete": False, "documents": {}, "episodes": []}
    try:
        data, entries, catalog, downloads, binding = load_manifest(manifest_path)
        total = data["total_episodes"]
        selected = data["selected_episodes"]
        report.update({"manifest_sha256": file_sha(manifest_path), "title": data["title"],
                       "series_id": data.get("series_id"), "total_episodes": total,
                       "selected_episodes": selected, "selected_count": len(selected),
                       "scope": "full_series" if len(selected) == total else "selected_episodes",
                       "download_binding": binding})
        raws, formatted = [], []
        for number in selected:
            entry, raw, corrected, layout = process_episode(number, entries.get(number), manifest_path.parent, catalog, downloads)
            artifacts = {}
            if raw is not None:
                episode_dir = out_dir / "episodes"
                for suffix, text in (("raw", raw), ("corrected", corrected), ("formatted", layout)):
                    artifacts[suffix] = str(write_versioned(episode_dir / (f"episode-{number:03d}." + suffix + ".txt"), text.encode("utf-8")))
            entry["artifacts"] = artifacts
            # Each audit is also versioned, so a newer manifest cannot make an
            # old review appear to belong to a newer raw transcript.
            audit = write_versioned(out_dir / "episodes" / f"episode-{number:03d}.audit.json", json_bytes(entry))
            entry["audit_file"] = str(audit)
            report["episodes"].append(entry)
            raws.append(raw)
            formatted.append(layout)
        report["missing_episodes"] = [item["episode"] for item in report["episodes"] if item["status"] == "missing"]
        report["failed_episodes"] = [item["episode"] for item in report["episodes"] if item["status"] == "failed"]
        report["pending_review_episodes"] = [item["episode"] for item in report["episodes"] if item["status"] == "pending_review"]
        report["recognition_complete"] = not report["missing_episodes"] and not report["failed_episodes"]
        report["complete"] = report["recognition_complete"] and not report["pending_review_episodes"]
        report["document_label"] = "校对版" if report["complete"] else "校对草稿"
        report["uncertain_count"] = sum(len(item["uncertain"]) for item in report["episodes"])
        if report["complete"] or not require_complete:
            stem = safe_name(data["title"])
            raw_body = render_document(data["title"], data.get("series_id"), total, selected, report["episodes"], raws, raw=True)
            body = render_document(data["title"], data.get("series_id"), total, selected, report["episodes"], formatted, complete=report["complete"])
            report["documents"]["raw_txt"] = str(write_versioned(out_dir / (stem + "_原始识别全文.txt"), raw_body.encode("utf-8")))
            report["documents"]["review_txt"] = str(write_versioned(out_dir / (stem + "_" + report["document_label"] + ".txt"), body.encode("utf-8")))
            if docx:
                report["documents"]["review_docx"] = str(write_versioned(out_dir / (stem + "_" + report["document_label"] + ".docx"), docx_bytes(body, data["title"])))
            report["document_status"] = "complete" if report["complete"] else "draft_with_placeholders_or_pending_review"
        else:
            report["document_status"] = "blocked_incomplete"
        code = 1 if report["failed_episodes"] else 0 if report["complete"] else 2
    except (OSError, ValueError, TypeError, KeyError) as exc:
        report.update({"complete": False, "document_status": "failed_validation_or_output", "error": str(exc)})
        code = 1
    report_path = write_versioned(out_dir / "report.json", json_bytes(report))
    return {"ok": code == 0, "complete": report["complete"], "report": str(report_path),
            "documents": report["documents"], "document_status": report["document_status"]}, code


def parser():
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--manifest", required=True, help="schema_version:1 的来源、原稿、校对与标点文件显式映射")
    result.add_argument("--out-dir", required=True, help="输出目录；已有不同内容另存 .v002 等版本")
    result.add_argument("--docx", action="store_true", help="另外输出 Word（需可选 python-docx）")
    result.add_argument("--require-complete", action="store_true", help="缺集、失败或待校对时只保存报告和每集审计，不输出汇总文档")
    return result


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    args = parser().parse_args(argv)
    try:
        result, code = build(args.manifest, args.out_dir, docx=args.docx, require_complete=args.require_complete)
    except (OSError, ValueError, TypeError, KeyError) as exc:
        result, code = {"ok": False, "error": str(exc)}, 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return code


if __name__ == "__main__":
    sys.exit(main())
